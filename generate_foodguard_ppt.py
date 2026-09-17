from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
NAVY = RGBColor(15, 36, 66)
BLUE = RGBColor(36, 99, 235)
TEAL = RGBColor(20, 143, 150)
LIGHT = RGBColor(239, 246, 255)
GRAY = RGBColor(88, 96, 107)
GREEN = RGBColor(41, 128, 92)
ORANGE = RGBColor(230, 126, 34)
RED = RGBColor(192, 57, 43)
WHITE = RGBColor(255, 255, 255)


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.title
    if title_box is None:
        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.0), Inches(0.7))
    title_box.text_frame.clear()
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.6), Inches(0.9), Inches(12.0), Inches(0.4))
        tf = sub.text_frame
        tf.text = subtitle
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.color.rgb = GRAY


def add_textbox(slide, left, top, width, height, text, font_size=18, color=NAVY, bold=False, fill=None, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 8
    tf.margin_right = 8
    tf.margin_top = 6
    tf.margin_bottom = 6
    if fill:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
        box.line.color.rgb = fill
    for idx, para in enumerate(tf.paragraphs):
        if idx > 0:
            para.level = 0
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_bullets(slide, left, top, width, height, items, font_size=18, color=NAVY):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 10
    tf.margin_right = 8
    tf.margin_top = 8
    tf.margin_bottom = 8
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(8)
    return box


def add_box(slide, left, top, width, height, text, fill_color, font_color=WHITE, font_size=20):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(font_size)
    p.font.bold = True
    p.font.color.rgb = font_color
    return shape


def add_arrow(slide, x1, y1, x2, y2, color=GRAY):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = Pt(2)
    return line


# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = LIGHT
# title text
box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.8), Inches(1.2))
text = box.text_frame
p = text.paragraphs[0]
p.text = 'FoodGuard'
p.font.size = Pt(30)
p.font.bold = True
p.font.color.rgb = NAVY
p2 = text.add_paragraph()
p2.text = '基於 MCP 與 RAG 的食品標示智慧合規判讀系統'
p2.font.size = Pt(24)
p2.font.bold = False
p2.font.color.rgb = BLUE

# subtitle
sub = slide.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(6.5), Inches(1.3))
sub_tf = sub.text_frame
sub_tf.word_wrap = True
for i, line in enumerate([
    '專題簡報：系統架構、設計理念與系統規格',
    '整合 RAG 檢索、MCP 工具協作、規則引擎與 Streamlit UI',
    '針對食品標示、營養宣稱、過敏原與法規依據進行判讀'
]):
    p = sub_tf.paragraphs[0] if i == 0 else sub_tf.add_paragraph()
    p.text = line
    p.font.size = Pt(21 if i == 0 else 16)
    p.font.color.rgb = GRAY if i > 0 else NAVY

# right panel
panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.2), Inches(1.1), Inches(4.2), Inches(3.8))
panel.fill.solid(); panel.fill.fore_color.rgb = WHITE; panel.line.color.rgb = BLUE
panel.line.width = Pt(2)
ptf = panel.text_frame
ptf.word_wrap = True
ptf.margin_left = 12
ptf.margin_right = 12
ptf.margin_top = 12
for idx, line in enumerate(['系統目標', '1. 提升食品標示判讀正確性', '2. 建立可追溯的法規依據', '3. 整合 AI + 規則引擎 + RAG', '4. 提供互動式 Web Demo']):
    p = ptf.paragraphs[0] if idx == 0 else ptf.add_paragraph()
    p.text = line
    p.font.size = Pt(18 if idx == 0 else 14)
    p.font.bold = idx == 0
    p.font.color.rgb = NAVY if idx == 0 else GRAY

# footer
footer = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.8), Inches(0.4))
ft = footer.text_frame
ft.text = 'FoodGuard｜專題 Demo｜2026'
ft.paragraphs[0].font.size = Pt(10)
ft.paragraphs[0].font.color.rgb = GRAY

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, '一、研究背景與系統目標', '食品標示合規檢查需兼顧法規依據、數值判斷與可解釋性')
add_bullets(slide, Inches(0.8), Inches(1.5), Inches(5.8), Inches(4.3), [
    '食品包裝資訊常涉及法規、營養宣稱與過敏原風險，判讀容易受主觀解讀影響。',
    '一般 AI 回答可能缺乏官方依據，導致判斷結果不具可驗證性。',
    '系統必須結合官方法規資料庫、規則引擎與追蹤來源能力。',
    '為了實務應用，需要提供結論、依據與建議，同時保留 UI 可操作性。'
], font_size=20)
add_bullets(slide, Inches(6.9), Inches(1.5), Inches(5.4), Inches(4.3), [
    '支援食品品名、成分、營養標示與營養宣稱輸入。',
    '利用 RAG 檢索 TFDA 官方文件與法規文本。',
    '結合 MCP 工具協作與 Python deterministic rule engine。',
    '輸出判讀結果、法規證據與可下載離線報告。'
], font_size=20)

# Slide 3: System architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, '二、系統架構', '以 User Interface、MCP Client、Server Tools 與 RAG 引擎構成分層式架構')

# boxes
left = Inches(0.7)
boxes = [
    ('Streamlit UI', Inches(0.7), Inches(1.8), Inches(1.6), Inches(1.2), BLUE),
    ('MCP Client', Inches(2.6), Inches(1.8), Inches(1.8), Inches(1.2), TEAL),
    ('MCP Server', Inches(4.8), Inches(1.8), Inches(1.8), Inches(1.2), GREEN),
    ('Rule Engine', Inches(7.0), Inches(1.8), Inches(1.8), Inches(1.2), ORANGE),
    ('RAG Retriever', Inches(9.2), Inches(1.8), Inches(1.9), Inches(1.2), RED),
    ('Official DOCs', Inches(11.4), Inches(1.8), Inches(1.2), Inches(1.2), NAVY),
    ('SQLite / Memory', Inches(5.1), Inches(3.8), Inches(2.5), Inches(1.0), BLUE),
    ('FAISS Index', Inches(8.3), Inches(3.8), Inches(2.2), Inches(1.0), TEAL),
    ('LLM / Summary', Inches(10.9), Inches(3.8), Inches(1.7), Inches(1.0), GREEN),
]
for text, x, y, w, h, color in boxes:
    add_box(slide, x, y, w, h, text, color, WHITE, 15)

# arrows
for a, b in [
    ((2.3, 2.4), (4.8, 2.4)),
    ((6.6, 2.4), (7.0, 2.4)),
    ((8.8, 2.4), (9.2, 2.4)),
    ((11.1, 2.4), (11.4, 2.4)),
    ((5.7, 3.0), (5.7, 3.8)),
    ((9.4, 3.0), (9.4, 3.8)),
    ((11.8, 3.0), (11.8, 3.8))
]:
    add_arrow(slide, a[0]*Inches(1), a[1]*Inches(1), b[0]*Inches(1), b[1]*Inches(1), color=GRAY)

# note
add_textbox(slide, Inches(0.8), Inches(5.5), Inches(11.8), Inches(1.2),
    '資料流：使用者輸入 → 解析產品資料 → MCP Client 調用工具 → RAG 檢索官方文件 → 規則引擎判讀 → LLM 整理結論與依據 → UI 顯示/下載報告',
    font_size=17, color=NAVY, bold=True)

# Slide 4: Design and module breakdown
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, '三、系統設計與模組分工', '設計重點：模組化、可追蹤、可驗證與可擴充')

add_bullets(slide, Inches(0.8), Inches(1.5), Inches(5.8), Inches(4.5), [
    'app.py：Streamlit 互動介面，負責產品輸入、顯示判讀結果與法規來源。',
    'foodguard/：負責輸入解析、產品資料結構化、規則判讀與上下文處理。',
    'backend/：處理 PDF 讀取、chunk 切分、embedding 與索引建立。',
    'rag/：建立 FAISS 向量儲存與檢索，保留 chunk metadata、source、page 與分數。'
], font_size=19)
add_bullets(slide, Inches(6.9), Inches(1.5), Inches(5.4), Inches(4.5), [
    'mcp_server.py：對外提供 search_food_regulation、check_allergens 等工具。',
    'mcp_client.py：負責與 MCP Server 通訊，統整 tool responses 與會話歷史。',
    'data/：保存 FAISS index、SQLite 記憶體、報告與訓練/判斷規則資料。',
    'documents/：官方法規資料來源，作為唯一判斷依據來源。'
], font_size=19)

# Slide 5 design principles
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, '四、設計原則', '系統設計遵守可驗證性與安全性')

items = [
    ('法規唯一來源', '法規依據僅取自 documents/，不自行生成法規內容，避免幻覺。'),
    ('可追溯證據', '每個判定保留文件名稱、頁碼、chunk ID、引用內容與相關度。'),
    ('規則優先', '數字比較與欄位判斷由 Python rule engine 執行，避免純 LLM 推論。'),
    ('失敗時保守', '若無足夠證據，統一顯示「目前知識庫找不到足夠依據」。'),
    ('MCP 協作', '工具分工明確，讓檢索、核對與摘要分離，提升系統可維護性。'),
    ('多層 fallback', '產品資料、DRIs、RAG 與 LLM 可作為保底，但不取代法規依據。')
]
for idx, (title, desc) in enumerate(items):
    y = Inches(1.5 + idx*0.9)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.8), Inches(0.7))
    box.fill.solid(); box.fill.fore_color.rgb = LIGHT; box.line.color.rgb = BLUE
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = f'{title}：{desc}'
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY

# Slide 6 system specification
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, '五、系統規格', '功能、環境與執行需求')

left = Inches(0.8)
# table-like boxes
def spec_block(x, y, w, h, title, lines, fill):
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    b.fill.solid(); b.fill.fore_color.rgb = fill; b.line.color.rgb = fill
    tf = b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = WHITE
    for line in lines:
        p = tf.add_paragraph(); p.text = line; p.font.size = Pt(12); p.font.color.rgb = WHITE; p.level = 0

spec_block(Inches(0.7), Inches(1.6), Inches(3.7), Inches(3.0), '開發環境', ['Windows / VS Code', 'Python 3.11', 'Streamlit', 'Official MCP Python SDK', 'OpenAI API / Ollama'], LIGHT)
spec_block(Inches(4.8), Inches(1.6), Inches(3.7), Inches(3.0), '核心資料', ['documents/ 法規文件', 'FAISS 向量索引', 'SQLite session memory', 'nutrition_claim_rules.json', 'food_composition.db'], BLUE)
spec_block(Inches(8.9), Inches(1.6), Inches(3.7), Inches(3.0), '核心能力', ['營養宣稱查核', '過敏原辨識', '法規 RAG 檢索', 'Web Search / 問答', '離線報告輸出'], TEAL)

add_bullets(slide, Inches(0.9), Inches(5.0), Inches(11.5), Inches(1.3), [
    '系統支援本機運行與 Web UI 展示，並可將結果保存至 SQLite 與 HTML 報告檔。',
    '資料處理流程以 deterministic rule engine 為主，並以 LLM 做整合式摘要與說明。'
], font_size=17)

# Slide 7 workflow
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, '六、系統工作流', '由輸入到判讀結論的完整執行流程')

# numbered steps with boxes
steps = [
    ('1. 輸入資料', '使用者輸入品名、成分、營養標示與宣稱。'),
    ('2. 解析產品', 'foodguard.parsing 正規化為 product_data 結構。'),
    ('3. 呼叫工具', 'MCP Client 透過 stdio 啟動 MCP Server。'),
    ('4. 檢索法規', 'RAG 搜尋相關文件與依據來源。'),
    ('5. 進行判讀', 'rule engine 比對欄位、門檻值與宣稱適用性。'),
    ('6. 輸出結果', 'UI 顯示結論、引證、建議與報告下載。')
]
for idx, (title, desc) in enumerate(steps):
    x = Inches(0.8 + (idx % 3) * 4.0)
    y = Inches(1.6 + (idx // 3) * 2.0)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.3), Inches(1.5))
    box.fill.solid(); box.fill.fore_color.rgb = LIGHT; box.line.color.rgb = BLUE
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(19); p.font.bold = True; p.font.color.rgb = NAVY
    p = tf.add_paragraph(); p.text = desc; p.font.size = Pt(12); p.font.color.rgb = GRAY

# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, '七、結論與未來發展', '系統可作為食品標示智慧檢查原型，兼顧可用性與法規可信度')
add_bullets(slide, Inches(0.8), Inches(1.7), Inches(5.6), Inches(3.8), [
    '本系統能將食品標示判讀從純文字問答，轉為具法規依據與可追溯的結構化判讀。',
    '透過 MCP 將檢索、規則與摘要模組拆分，有助於維護與延展。',
    'Faiss + RAG + 公式規則結合，使判讀更具可信度與一致性。',
    '未來可以擴大資料涵蓋範圍，整合更多官方食品法規與數據庫。'
], font_size=20)
add_bullets(slide, Inches(6.8), Inches(1.7), Inches(5.5), Inches(3.8), [
    '可加入更多食品類別與營養項目判讀邏輯。',
    '可整合雲端部署與使用者帳號管理。',
    '可強化 Web Search + FAQ + 風險提醒功能。',
    '可發展為實際合規審核工具或校內研究展示平台。'
], font_size=20)

# final summary footer
slide.shapes.add_textbox(Inches(0.8), Inches(6.7), Inches(11.7), Inches(0.3)).text_frame.text = 'FoodGuard：結合 MCP、RAG、規則引擎與 Streamlit 的食品標示智慧判讀架構'

ppt_file = 'FoodGuard_系統架構與規格簡報.pptx'
prs.save(ppt_file)
print(f'PPT created: {ppt_file}')
